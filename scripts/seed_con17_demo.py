#!/usr/bin/env python3
"""Create a disposable CON-17 browser fixture: eight ended synthetic meetings.

The default database is intentionally separate from Convene's normal data/convene.db. A manifest records every
created id; scripts/remove_con17_demo.py requires that manifest before removing anything.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from server import registry
from server.db import Database
from server.ids import new_id
from server.repositories import utterances
from server.repositories.models import Utterance
from server.timeutil import utc_now
from docx import Document

DEFAULT_DB = ROOT / "data" / "con17-policy-demo.db"
DEFAULT_MANIFEST = ROOT / "data" / "con17-policy-demo-manifest.json"

FIXTURES = (
    ("Sprint planning", "Priya", "The beta launch is planned for 15 October."),
    ("Travel policy review", "Marcus", "The travel budget needs director approval."),
    ("Design review", "Dana", "The dashboard needs a visible policy source label."),
    ("Vendor selection", "Asha", "The procurement rule requires two quotations."),
    ("Security check-in", "Ben", "Password rotations happen every ninety days."),
    ("Hiring plan", "Chitra", "The intern starts on the first Monday in November."),
    ("Operations sync", "Dev", "Incident reports need an owner before close."),
    ("Quarterly retrospective", "Eve", "The team will revisit travel spend next quarter."),
)


POLICIES = {
    "travel-policy-v1.docx": "Travel policy version 1. Flights need director approval and itemised receipts.",
    "travel-policy-v2.docx": "Travel policy version 2. Flights over five hundred dollars need director approval and itemised receipts.",
    "procurement-policy.docx": "Procurement policy. Purchases above ten thousand require two written quotations.",
}


def _write_policy_inputs(directory: Path) -> list[str]:
    directory.mkdir(parents=True, exist_ok=True)
    paths=[]
    for name, text in POLICIES.items():
        document=Document(); document.add_paragraph(text)
        table=document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text="Fixture"
        table.rows[0].cells[1].text="Synthetic only"
        path=directory/name; document.save(path); paths.append(str(path.resolve()))
    return paths


def _write_demo_config(directory: Path, database: Path) -> Path:
    """Copy the normal config so model pins stay identical, replacing only the database path."""
    source=(ROOT / "config" / "convene.toml").read_text(encoding="utf-8")
    replacement=f'database = "{database.resolve()}"'
    config=re.sub(r'^database\s*=\s*"[^"]+"', replacement, source, count=1, flags=re.MULTILINE)
    output=directory / "con17-policy-demo.toml"
    output.write_text(config, encoding="utf-8")
    return output


def seed(path: Path, manifest_path: Path) -> dict:
    if manifest_path.exists():
        raise SystemExit(f"manifest already exists: {manifest_path}; run remove_con17_demo.py first")
    database = Database.open(path)
    meetings = []
    try:
        for title, speaker, text in FIXTURES:
            with database.transaction() as tx:
                meeting = registry.create_meeting(tx, title)
                registration = registry.register_device(tx, meeting.meeting_id, new_id(), speaker)
                now = utc_now()
                utterance = utterances.insert(tx.conn, Utterance(
                    new_id(), meeting.meeting_id, registration.device.device_id,
                    registration.participants[0].participant_id, text, now, now, .95, "device", .95, now))
                registry.end_meeting(tx, meeting.meeting_id)
                meetings.append({"meeting_id": meeting.meeting_id, "utterance_id": utterance.utterance_id})
    finally:
        database.close()
    inputs = _write_policy_inputs(manifest_path.parent / "con17-policy-demo-inputs")
    config = _write_demo_config(manifest_path.parent, path)
    manifest = {"schema": 1, "database": str(path.resolve()), "meetings": meetings, "files": inputs,
                "config": str(config.resolve())}
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", type=Path, default=DEFAULT_DB)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    args = parser.parse_args()
    result = seed(args.database, args.manifest)
    print(f"Created {len(result['meetings'])} synthetic ended meetings in {args.database}")
    print(f"Manifest: {args.manifest}")
    print(f"Synthetic policy uploads: {args.manifest.parent / 'con17-policy-demo-inputs'}")
    print(f"Start with: .venv/bin/python -m server.app --config {result['config']} --cert <cert> --key <key>")


if __name__ == "__main__":
    main()
