"""Pure fixed-template DOCX renderer.  It consumes rows/views only and performs no I/O."""
from __future__ import annotations

from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt


def _utc_date(value: str | None) -> str:
    if not value:
        return "Date unavailable"
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone.utc).strftime("%Y-%m-%d UTC")


def elapsed(start: str | None, value: str) -> str:
    if not start:
        return "00:00:00"
    seconds = max(0, int((datetime.fromisoformat(value.replace("Z", "+00:00")) -
                          datetime.fromisoformat(start.replace("Z", "+00:00"))).total_seconds()))
    return f"{seconds // 3600:02d}:{seconds // 60 % 60:02d}:{seconds % 60:02d}"


def render(meeting, participants, summary, action_items, utterances) -> Document:
    """Build the fixed minutes template. ``utterances`` are ordered server-computed views."""
    document = Document()
    core = document.core_properties
    # Semantic document XML is stable across equivalent renders; ZIP container timestamps are not.
    core.created = core.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.add_heading(meeting.title or f"Meeting {meeting.meeting_id}", 0)
    document.add_paragraph(_utc_date(meeting.started_at or meeting.created_at))
    names = ", ".join(person.display_name for person in participants) or "No participants registered"
    document.add_paragraph(f"Participants: {names}")

    document.add_heading("Summary", level=1)
    for paragraph in summary.summary_text.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())

    document.add_heading("Action items", level=1)
    if not action_items:
        document.add_paragraph("No action items")
    else:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ("Item", "Owner", "Status")):
            cell.text = label
        for item in action_items:
            cells = table.add_row().cells
            cells[0].text = item["text"]
            cells[1].text = item.get("owner_display_name") or "Unassigned"
            cells[2].text = item["status"]

    document.add_heading("Transcript appendix", level=1)
    start = meeting.started_at
    for item in utterances:
        paragraph = document.add_paragraph()
        if item["low_confidence"]:
            paragraph.style = document.styles["Normal"]
            paragraph.paragraph_format.left_indent = Pt(8)
        prefix = paragraph.add_run(f"[{elapsed(start, item['t_start'])}] {item['speaker_label']}: ")
        prefix.bold = True
        text = paragraph.add_run(item["text"])
        if item["low_confidence"]:
            text.italic = True
            paragraph.add_run(" [needs review]").italic = True
    return document
