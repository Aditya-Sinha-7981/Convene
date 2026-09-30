"""Pure fixed-template DOCX renderer for the periodic report (CON-18). No I/O, no database, no clock.

A second template on the minutes renderer's conventions (docs/export.md, ADR-12): fixed core timestamps, UTC dates,
equivalent document XML for equal input. ``data`` is ``server.reports.service.build``'s output; ``generated_at`` is
passed in so the output never depends on when it runs.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from docx import Document

from .docx_renderer import _utc_date, elapsed

SUMMARY_STATES = {
    "none": "No summary",
    "pending": "No summary (a summary was still being written when this report was generated)",
}


def _table(document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value


def _range_text(data: dict) -> str:
    return f"{_utc_date(data['from'])} to {_utc_date(data['to'])}"


def _figure(document, heading: str, definition: str, items: list[dict]) -> None:
    document.add_heading(f"{heading}: {len(items)}", level=2)
    document.add_paragraph(definition)
    if not items:
        document.add_paragraph("None")
        return
    _table(document, ("Item", "Owner", "Due", "Status", "Meeting"),
           [(i["text"], i["owner"], i["due_date"] or "—", i["status"],
             f"{i['meeting_title']} ({_utc_date(i['meeting_date'])})") for i in items])


def render_report(data: dict, generated_at: str) -> Document:
    document = Document()
    core = document.core_properties
    core.created = core.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    meetings = data["meetings"]

    document.add_heading("Convene report", 0)
    document.add_paragraph(f"Period: {_range_text(data)} (dates are UTC)")
    document.add_paragraph(f"Generated: {_utc_date(generated_at)}")
    counts = f"Meetings: {len(meetings)} ended, {data['with_summary_count']} with a summary"
    if data["excluded_not_ended_count"]:
        counts += f"; {data['excluded_not_ended_count']} not yet ended, not included"
    document.add_paragraph(counts)

    document.add_heading("Action items", level=1)
    _figure(document, "Opened in this period", "Items from the current summary of each meeting in this report.",
            data["opened"])
    _figure(document, "Closed in this period (done)",
            "Items from any meeting whose last status change in this period marked them done.", data["closed_done"])
    _figure(document, "Cancelled in this period",
            "Items from any meeting whose last status change in this period marked them cancelled.",
            data["closed_cancelled"])
    _figure(document, "Open now", "Items from this report's meetings still open when the report was generated "
                                  "(not as of the end of the period).", data["open_now"])

    document.add_heading("Meetings", level=1)
    for meeting in meetings:
        document.add_heading(f"{meeting['title']} — {_utc_date(meeting['started_at'] or meeting['created_at'])}",
                             level=2)
        if meeting["started_at"] and meeting["ended_at"]:
            document.add_paragraph(f"Duration: {elapsed(meeting['started_at'], meeting['ended_at'])}")
        names = ", ".join(meeting["participants"]) or "No participants registered"
        document.add_paragraph(f"Participants: {names}")
        document.add_heading("Summary", level=3)
        if meeting["summary_state"] == "ready":
            if meeting["stale"]:
                document.add_paragraph("Note: the transcript changed after this summary was written, so it may be "
                                       "out of date.").runs[0].italic = True
            for paragraph in (meeting["summary_text"] or "").split("\n\n"):
                if paragraph.strip():
                    document.add_paragraph(paragraph.strip())
        elif meeting["summary_state"] == "failed":
            document.add_paragraph(f"Summary failed: {meeting['failure_code']}")
        else:
            document.add_paragraph(SUMMARY_STATES[meeting["summary_state"]])
        document.add_heading("Action items", level=3)
        if not meeting["action_items"]:
            document.add_paragraph("No action items")
        else:
            _table(document, ("Item", "Owner", "Due", "Status"),
                   [(i["text"], i["owner"], i["due_date"] or "—", i["status"]) for i in meeting["action_items"]])
    return document


SELECTION_STATES = {
    "none": "No summary yet.",
    "pending": "A summary was still being written when this file was made.",
}


def render_summaries(sections: list[dict], generated_at: str) -> Document:
    """Several meetings' summaries in one file (history page): each title and date, participants, then its summary.

    ``sections`` are ``server.reports.service.selected_sections`` output, oldest first. No action items or transcript:
    it is the "Summaries" tab as a document. Same conventions as the report: fixed core timestamps, UTC dates.
    """
    document = Document()
    core = document.core_properties
    core.created = core.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.add_heading("Convene meeting summaries", 0)
    document.add_paragraph(f"Generated: {_utc_date(generated_at)}")
    with_summary = sum(1 for meeting in sections if meeting["summary_state"] == "ready")
    count = "1 meeting" if len(sections) == 1 else f"{len(sections)} meetings"
    document.add_paragraph(f"{count}, {with_summary} with a summary, oldest first (dates are UTC)")
    for meeting in sections:
        document.add_heading(f"{meeting['title']} — {_utc_date(meeting['started_at'] or meeting['created_at'])}", level=1)
        names = ", ".join(meeting["participants"]) or "No participants registered"
        document.add_paragraph(f"Participants: {names}")
        if meeting["summary_state"] == "ready":
            if meeting["stale"]:
                document.add_paragraph("Note: the transcript changed after this summary was written, so it may be "
                                       "out of date.").runs[0].italic = True
            for paragraph in (meeting["summary_text"] or "").split("\n\n"):
                if paragraph.strip():
                    document.add_paragraph(paragraph.strip())
        elif meeting["summary_state"] == "failed":
            document.add_paragraph(f"No summary: the last attempt failed ({meeting['failure_code']}).")
        else:
            document.add_paragraph(SELECTION_STATES[meeting["summary_state"]])
    return document


def render_summaries_bytes(sections: list[dict], generated_at: str) -> bytes:
    buffer = io.BytesIO()
    render_summaries(sections, generated_at).save(buffer)
    return buffer.getvalue()


def render_report_bytes(data: dict, generated_at: str) -> bytes:
    buffer = io.BytesIO()
    render_report(data, generated_at).save(buffer)
    return buffer.getvalue()
