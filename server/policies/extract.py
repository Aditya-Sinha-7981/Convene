"""Local, deliberately non-OCR extraction for policy uploads."""
from __future__ import annotations
import io
import re
import zipfile
from docx import Document
from pypdf import PdfReader

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

class ExtractionError(Exception):
    def __init__(self, code, message): self.code, self.message = code, message

def media_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"): return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "[Content_Types].xml" in archive.namelist() and "word/document.xml" in archive.namelist(): return DOCX
        except zipfile.BadZipFile: pass
    raise ExtractionError("unsupported_format", "only PDF and DOCX policy files are supported")

def _normalise(text: str) -> str:
    paragraphs = [re.sub(r"[ \t]+", " ", p).strip() for p in re.split(r"\n{2,}", text.replace("\r", "\n"))]
    return "\n\n".join(p for p in paragraphs if p)

def extract(data: bytes, kind: str) -> str:
    try:
        if kind == PDF:
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif kind == DOCX:
            document = Document(io.BytesIO(data))
            bits = [p.text for p in document.paragraphs]
            bits.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            text = "\n\n".join(bits)
        else: raise ExtractionError("unsupported_format", "only PDF and DOCX policy files are supported")
    except ExtractionError: raise
    except Exception as exc: raise ExtractionError("extract_failed", "the document could not be read locally") from exc
    text = _normalise(text)
    if not text: raise ExtractionError("no_text_layer", "no extractable text was found; scanned PDFs need OCR, which Convene does not perform")
    return text
