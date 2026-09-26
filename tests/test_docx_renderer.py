from types import SimpleNamespace

from docx import Document

from server.export.docx_renderer import render


def _meeting():
    return SimpleNamespace(meeting_id="00000000-0000-4000-8000-000000000000", title="Planning ✨",
                           created_at="2026-09-26T10:00:00.000Z", started_at="2026-09-26T10:00:00.000Z")


def _summary():
    return SimpleNamespace(summary_text="The team planned the launch.\n\nSam owns the notes.")


def test_fixed_template_contains_stored_content_in_order(tmp_path):
    doc = render(_meeting(), [SimpleNamespace(display_name="Ada 😊")], _summary(),
                 [{"text": "Write notes", "owner_display_name": None, "status": "open"}],
                 [{"t_start": "2026-09-26T10:00:05.000Z", "speaker_label": "Ada 😊", "text": "Hello",
                   "low_confidence": False},
                  {"t_start": "2026-09-26T10:00:06.000Z", "speaker_label": "Speaker on Phone 2", "text": "Maybe",
                   "low_confidence": True}])
    path = tmp_path / "minutes.docx"
    doc.save(path)
    opened = Document(path)
    text = "\n".join(p.text for p in opened.paragraphs)
    assert text.index("Planning ✨") < text.index("Summary") < text.index("Action items") < text.index("Transcript appendix")
    assert "Ada 😊" in text and "[needs review]" in text
    assert [cell.text for cell in opened.tables[0].rows[1].cells] == ["Write notes", "Unassigned", "open"]


def test_empty_actions_and_low_confidence_formatting():
    doc = render(_meeting(), [], _summary(), [], [{"t_start": "2026-09-26T10:00:01.000Z",
                 "speaker_label": "Speaker on Phone 1", "text": "Uncertain", "low_confidence": True}])
    assert any(p.text == "No action items" for p in doc.paragraphs)
    line = next(p for p in doc.paragraphs if "Uncertain" in p.text)
    assert any(run.italic for run in line.runs)
