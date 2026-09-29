"""Local-only policy extraction and content-type validation."""
import io

import pytest
from docx import Document

from server.policies.extract import DOCX, PDF, ExtractionError, extract, media_type


def _pdf(text: str = "") -> bytes:
    # Small deliberately generated fixture; offsets are calculated so pypdf reads it like a normal text PDF.
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET" if text else ""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
               f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    body = "%PDF-1.4\n"; offsets=[]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(body.encode())); body += f"{n} 0 obj\n{obj}\nendobj\n"
    xref=len(body.encode()); body += "xref\n0 6\n0000000000 65535 f \n" + "".join(f"{offset:010} 00000 n \n" for offset in offsets)
    return (body + f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n").encode()


def test_text_pdf_extracts_and_a_scanned_like_pdf_fails_clearly():
    assert media_type(_pdf("Travel approval is required")) == PDF
    assert "Travel approval" in extract(_pdf("Travel approval is required"), PDF)
    with pytest.raises(ExtractionError, match="no extractable text") as failed:
        extract(_pdf(), PDF)
    assert failed.value.code == "no_text_layer"


def test_docx_extracts_paragraphs_and_table_cells():
    document=Document(); document.add_paragraph("Paragraph policy text")
    row=document.add_table(rows=1, cols=2).rows[0]; row.cells[0].text="Expense"; row.cells[1].text="Receipt"
    stream=io.BytesIO(); document.save(stream)
    assert media_type(stream.getvalue()) == DOCX
    text=extract(stream.getvalue(), DOCX)
    assert "Paragraph policy text" in text and "Expense | Receipt" in text


def test_spoofed_extension_cannot_change_magic_validation():
    with pytest.raises(ExtractionError) as failed:
        media_type(b"not a policy")
    assert failed.value.code == "unsupported_format"
