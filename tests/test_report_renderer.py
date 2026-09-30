"""CON-18: the pure report renderer (fixed template, fixed core timestamps, UTC dates, equivalent XML)."""
from datetime import datetime, timezone
from io import BytesIO

from docx import Document

from server.export.report_renderer import render_report, render_report_bytes

MEETING_ID = "0d4f6a52-7c1b-4e7a-b0a3-51e1f4c2a9d8"


def item(text, status="open", owner="Sam", due=None):
    return {"action_item_id": "f0a3c5e7-1b92-4d68-8c4a-3e7d9b1f2a60", "text": text, "owner": owner, "status": status,
            "due_date": due, "meeting_id": MEETING_ID, "meeting_title": "Sprint",
            "meeting_date": "2026-09-21T23:30:00.000+05:30"}


def meeting(**overrides):
    value = {"meeting_id": MEETING_ID, "title": "Sprint", "created_at": "2026-09-21T11:30:00.000Z",
             "started_at": "2026-09-21T11:31:00.000Z", "ended_at": "2026-09-21T12:01:30.000Z",
             "participants": ["Priya", "Sam"], "summary_state": "ready",
             "summary_text": "We planned.\n\nSam writes notes.", "failure_code": None, "stale": False,
             "action_items": [item("Write notes", due="2026-09-25")]}
    value.update(overrides)
    return value


def data(meetings=None, **overrides):
    meetings = [meeting()] if meetings is None else meetings
    opened = [i for m in meetings for i in m["action_items"]]
    value = {"from": "2026-09-01T00:00:00.000Z", "to": "2026-09-30T23:59:59.999Z", "excluded_not_ended_count": 0,
             "with_summary_count": sum(m["summary_state"] == "ready" for m in meetings), "meetings": meetings,
             "opened": opened, "closed_done": [], "closed_cancelled": [],
             "open_now": [i for i in opened if i["status"] == "open"]}
    value.update(overrides)
    return value


def text(document):
    return "\n".join(p.text for p in document.paragraphs)


def xml(document):
    return document.element.xml


def test_same_input_gives_equivalent_xml_and_fixed_core_timestamps():
    a, b = render_report(data(), "2026-09-30T08:00:00.000Z"), render_report(data(), "2026-09-30T08:00:00.000Z")
    assert xml(a) == xml(b)
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    reopened = Document(BytesIO(render_report_bytes(data(), "2026-09-30T08:00:00.000Z")))
    assert reopened.core_properties.created == fixed and reopened.core_properties.modified == fixed


def test_headings_figures_and_tables_in_order():
    doc = render_report(data(closed_done=[item("Old thing", status="done")]), "2026-09-30T08:00:00.000Z")
    body = text(doc)
    order = ["Convene report", "Period: 2026-09-01 UTC to 2026-09-30 UTC", "Generated: 2026-09-30 UTC",
             "Opened in this period: 1", "Closed in this period (done): 1", "Cancelled in this period: 0",
             "Open now: 1", "\nMeetings\n", "Sprint — 2026-09-21 UTC", "Duration: 00:30:30", "Participants: Priya, Sam",
             "We planned.", "Sam writes notes."]
    positions = [body.index(part) for part in order]
    assert positions == sorted(positions)
    tables = doc.tables
    assert [c.text for c in tables[0].rows[0].cells] == ["Item", "Owner", "Due", "Status", "Meeting"]
    assert [c.text for c in tables[0].rows[1].cells] == ["Write notes", "Sam", "2026-09-25", "open", "Sprint (2026-09-21 UTC)"]
    assert [c.text for c in tables[-1].rows[1].cells] == ["Write notes", "Sam", "2026-09-25", "open"]


def test_dates_are_utc():
    """A meeting at 23:30 in India is 18:00 UTC the same day; one at 02:00 +05:30 is the previous UTC day."""
    doc = render_report(data([meeting(started_at="2026-09-22T02:00:00.000+05:30")]), "2026-09-30T08:00:00.000Z")
    assert "Sprint — 2026-09-21 UTC" in text(doc)


def test_summary_states_are_explicit_and_stale_is_noted():
    meetings = [meeting(title="Stale", stale=True), meeting(title="Failed", summary_state="failed", summary_text=None,
                failure_code="summary_invalid_output", action_items=[]),
                meeting(title="None", summary_state="none", summary_text=None, action_items=[]),
                meeting(title="Pending", summary_state="pending", summary_text=None, action_items=[])]
    body = text(render_report(data(meetings), "2026-09-30T08:00:00.000Z"))
    assert "may be out of date" in body
    assert "Summary failed: summary_invalid_output" in body
    assert "\nNo summary\n" in body and "still being written" in body
    assert body.count("No action items") == 3


def test_zero_meetings_zero_items_and_long_summaries():
    empty = render_report(data([]), "2026-09-30T08:00:00.000Z")
    assert "Meetings: 0 ended, 0 with a summary" in text(empty) and text(empty).count("None") == 4
    long = "Word " * 5000 + "\n\n" + "More " * 5000
    doc = render_report(data([meeting(summary_text=long, action_items=[])]), "2026-09-30T08:00:00.000Z")
    assert text(doc).count("Word") == 5000 and text(doc).count("More") == 5000


def test_excluded_meetings_are_mentioned():
    body = text(render_report(data(excluded_not_ended_count=2), "2026-09-30T08:00:00.000Z"))
    assert "2 not yet ended, not included" in body
