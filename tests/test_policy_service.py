"""CON-17 append-only local policy ingestion (no HTTP server required)."""
import io
from docx import Document
import pytest

from server.config import Settings
from server.db import Database
from server.policies.service import PolicyService
from server.rag.embedding import FakeEmbeddingAdapter
from server.repositories import policies


def docx_bytes(text):
    document=Document(); document.add_paragraph(text); document.add_table(rows=1, cols=2).rows[0].cells[0].text="Owner"
    stream=io.BytesIO(); document.save(stream); return stream.getvalue()


@pytest.mark.asyncio
async def test_docx_upload_is_append_only_and_indexed(tmp_path):
    settings=Settings(root=tmp_path,database_path=tmp_path/"db.sqlite",exports_dir=tmp_path/"exports",policies_dir=tmp_path/"policies")
    db=Database.open(settings.database_path); service=PolicyService(db,settings,FakeEmbeddingAdapter(dimension=384))
    first_bytes=docx_bytes("Flights need manager approval.")
    first=await service.upload("Travel policy",["Travel"],"first.docx",first_bytes)
    version1=first["version"]; assert version1["status"]=="pending"
    await __import__('asyncio').gather(*service._tasks)
    ready1=policies.version(db.conn,version1["policy_version_id"]); assert ready1.status=="ready" and "Flights" in ready1.extracted_text
    second=await service.upload(None,[],"second.docx",docx_bytes("Flights need director approval."),first["policy"]["policy_id"])
    await __import__('asyncio').gather(*service._tasks)
    versions=policies.versions(db.conn,first["policy"]["policy_id"])
    assert [row.version_number for row in versions]==[2,1]
    assert (settings.policies_dir/ready1.storage_path).read_bytes() == first_bytes
    assert policies.current_versions(db.conn)[0].version_number==2
    db.close()


@pytest.mark.asyncio
async def test_empty_document_is_retained_as_a_visible_no_text_failure(tmp_path):
    settings=Settings(root=tmp_path,database_path=tmp_path/"db.sqlite",exports_dir=tmp_path/"exports",policies_dir=tmp_path/"policies")
    db=Database.open(settings.database_path); service=PolicyService(db,settings,FakeEmbeddingAdapter(dimension=384))
    empty=Document(); stream=io.BytesIO(); empty.save(stream)
    result=await service.upload("Scanned",[],"scan.docx",stream.getvalue())
    await __import__('asyncio').gather(*service._tasks)
    version=policies.version(db.conn,result["version"]["policy_version_id"])
    assert version.status=="failed" and version.error_code=="no_text_layer"
    assert (settings.policies_dir/version.storage_path).is_file()
    db.close()
