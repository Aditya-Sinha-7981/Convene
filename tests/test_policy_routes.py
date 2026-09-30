"""CON-17 policy HTTP error mapping: bad uploads are 400s and retry only accepts failed versions."""
import asyncio
import io

import pytest
from aiohttp import ClientSession, FormData
from docx import Document

from server.ids import new_id
from server.rag.embedding import FakeEmbeddingAdapter
from server.repositories import policies
from tests.support.server import start_server


def docx_bytes(text):
    document = Document(); document.add_paragraph(text)
    stream = io.BytesIO(); document.save(stream); return stream.getvalue()


def upload(data, filename="policy.docx", title=None):
    form = FormData()
    if title is not None:
        form.add_field("title", title)
    form.add_field("file", data, filename=filename)
    return form


@pytest.fixture
async def server(settings):
    """The real app with a fake local embedding model, so a text upload becomes ready."""
    running = await start_server(settings, embedding_adapter=FakeEmbeddingAdapter(dimension=384))
    yield running
    await running.stop()


@pytest.fixture
async def http():
    async with ClientSession() as session:
        yield session


async def create(server, http, data, title="Travel policy"):
    async with http.post(server.base_url + "/api/policies", data=upload(data, title=title)) as response:
        assert response.status == 202
        body = await response.json()
    await asyncio.gather(*server.runtime.policies._tasks, return_exceptions=True)
    return body


@pytest.mark.asyncio
async def test_new_version_rejects_unsupported_and_empty_files_as_400(server, http):
    created = await create(server, http, docx_bytes("Flights need approval."))
    url = server.base_url + f"/api/policies/{created['policy']['policy_id']}/versions"
    async with http.post(url, data=upload(b"not a policy", "spoof.pdf")) as response:
        assert response.status == 400 and (await response.json())["error"]["code"] == "unsupported_format"
    async with http.post(url, data=upload(b"", "empty.docx")) as response:
        assert response.status == 400 and (await response.json())["error"]["code"] == "invalid_request"
    async with http.post(server.base_url + f"/api/policies/{new_id()}/versions",
                         data=upload(docx_bytes("x"))) as response:
        assert response.status == 404 and (await response.json())["error"]["code"] == "policy_not_found"


@pytest.mark.asyncio
async def test_retry_accepts_only_a_failed_version_of_the_named_policy(server, http):
    empty = Document(); stream = io.BytesIO(); empty.save(stream)
    failed = await create(server, http, stream.getvalue(), title="Scanned")
    ready = await create(server, http, docx_bytes("Flights need approval."))
    db = server.runtime.db

    ready_version = ready["version"]["policy_version_id"]
    async with http.post(server.base_url + f"/api/policies/{ready['policy']['policy_id']}/versions/{ready_version}/retry",
                         json={}) as response:
        assert response.status == 409 and (await response.json())["error"]["code"] == "policy_version_not_failed"
    assert policies.version(db.conn, ready_version).status == "ready"  # untouched

    failed_version = failed["version"]["policy_version_id"]
    assert policies.version(db.conn, failed_version).status == "failed"
    # The version exists but under another policy: a 404, and it is not requeued.
    async with http.post(server.base_url + f"/api/policies/{ready['policy']['policy_id']}/versions/{failed_version}/retry",
                         json={}) as response:
        assert response.status == 404 and (await response.json())["error"]["code"] == "policy_not_found"
    assert policies.version(db.conn, failed_version).status == "failed"

    async with http.post(server.base_url + f"/api/policies/{failed['policy']['policy_id']}/versions/{failed_version}/retry",
                         json={}) as response:
        assert response.status == 202
    await asyncio.gather(*server.runtime.policies._tasks, return_exceptions=True)
    assert policies.version(db.conn, failed_version).error_code == "no_text_layer"  # reprocessed, still no text
