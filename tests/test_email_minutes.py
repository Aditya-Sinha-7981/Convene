"""ADR-33: an optional email at join, and emailing the DOCX minutes through Resend on request.

The Resend API is never called: the connector is exercised with a fake ``urlopen`` and the API tests use a fake
mailer. That a real message arrives with the attachment is a manual check (logs/email-report.md).
"""
import base64
import io
import json
import urllib.error
from dataclasses import dataclass, field

import pytest
from aiohttp import ClientSession
from docx import Document

from server import registry
from server.errors import ValidationError
from server.ids import new_id
from server.mail import message
from server.mail.resend import (ENDPOINT, Attachment, MailConfigError, MailError, ResendMailer,
                                mailer_from_environment)
from server.export.docx_renderer import render
from server.mail.service import mask
from server.repositories import participant_emails
from tests.support.reports import meeting_at
from tests.support.server import settings_in, start_server

KEY = "re_secret_key_123"


def make_meeting(db, title="Sprint planning"):
    with db.transaction() as tx:
        return registry.create_meeting(tx, title)


def register(db, meeting_id, name="Priya", device_id=None, **kwargs):
    with db.transaction() as tx:
        return registry.register_device(tx, meeting_id, device_id or new_id(), name, **kwargs)


# --- registration ----------------------------------------------------------------------------


def test_email_is_stored_apart_from_the_participant(db):
    meeting = make_meeting(db)
    person = register(db, meeting.meeting_id, email="  priya@example.com ").participants[0]
    assert participant_emails.get(db.conn, person.participant_id) == "priya@example.com"
    assert not hasattr(person, "email")
    audit = db.conn.execute("SELECT group_concat(payload) FROM AuditEvent").fetchone()[0]
    assert "priya@example.com" not in audit


@pytest.mark.parametrize("bad", ["priya", "priya@example", "a b@example.com", "x@y.z<script>", "a@" + "b" * 260 + ".com", 7])
def test_invalid_email_is_rejected_and_nothing_is_written(db, bad):
    meeting = make_meeting(db)
    with pytest.raises(ValidationError):
        register(db, meeting.meeting_id, email=bad)
    assert db.conn.execute("SELECT COUNT(*) FROM Device").fetchone()[0] == 0


def test_blank_email_means_none_and_shared_devices_cannot_give_one(db):
    meeting = make_meeting(db)
    person = register(db, meeting.meeting_id, email="   ").participants[0]
    assert participant_emails.get(db.conn, person.participant_id) is None
    with pytest.raises(ValidationError):
        with db.transaction() as tx:
            registry.register_device(tx, meeting.meeting_id, new_id(), None, True, 2, email="a@b.co")


def test_a_rejoin_can_add_or_change_the_address_and_a_change_resets_last_sent(db):
    meeting = make_meeting(db)
    device_id = new_id()
    person = register(db, meeting.meeting_id, device_id=device_id).participants[0]
    assert participant_emails.get(db.conn, person.participant_id) is None
    replay = register(db, meeting.meeting_id, device_id=device_id, email="first@example.com")
    assert replay.created is False
    participant_emails.mark_sent(db.conn, person.participant_id, "2026-09-30T10:00:00.000Z")
    register(db, meeting.meeting_id, device_id=device_id, email="first@example.com")  # same address: kept as sent
    assert participant_emails.list_for_meeting(db.conn, meeting.meeting_id)[0]["last_sent_at"] is not None
    register(db, meeting.meeting_id, device_id=device_id, email="second@example.com")
    row = participant_emails.list_for_meeting(db.conn, meeting.meeting_id)[0]
    assert (row["email"], row["last_sent_at"]) == ("second@example.com", None)
    register(db, meeting.meeting_id, device_id=device_id)  # a replay without one keeps it
    assert participant_emails.get(db.conn, person.participant_id) == "second@example.com"


def test_erasing_the_meeting_erases_the_addresses(db):
    meeting = make_meeting(db)
    register(db, meeting.meeting_id, email="priya@example.com")
    with db.transaction() as tx:
        registry.erase_meeting(tx, meeting.meeting_id)
    assert db.conn.execute("SELECT COUNT(*) FROM ParticipantEmail").fetchone()[0] == 0


# --- the Resend connector ----------------------------------------------------------------------


class FakeResponse(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def http_error(code, body):
    return urllib.error.HTTPError(ENDPOINT, code, "err", {}, io.BytesIO(json.dumps(body).encode()))


def test_connector_sends_one_message_with_the_attachment():
    seen = []

    def opener(request, timeout):
        seen.append(request)
        return FakeResponse(b'{"id": "msg_1"}')

    mailer = ResendMailer(KEY, "Convene <minutes@example.com>", opener=opener)
    assert mailer.send("a@b.co", "Subj", "<p>h</p>", "t", Attachment("m.docx", b"PK\x03\x04")) == "msg_1"
    request = seen[0]
    body = json.loads(request.data)
    assert request.full_url == ENDPOINT and request.get_method() == "POST"
    assert request.get_header("Authorization") == f"Bearer {KEY}" and request.get_header("User-agent")
    assert body["to"] == ["a@b.co"] and body["from"] == "Convene <minutes@example.com>"
    assert body["attachments"] == [{"filename": "m.docx", "content": base64.b64encode(b"PK\x03\x04").decode()}]
    assert KEY not in repr(mailer)


def test_connector_retries_a_rate_limit_once_then_reports_without_the_key():
    calls = []

    def opener(request, timeout):
        calls.append(1)
        raise http_error(429 if len(calls) == 1 else 422, {"message": "Invalid `to` field."})

    mailer = ResendMailer(KEY, "minutes@example.com", opener=opener, sleep=lambda s: None)
    with pytest.raises(MailError) as caught:
        mailer.send("a@b.co", "s", "h", "t")
    assert len(calls) == 2 and "422" in str(caught.value) and "Invalid" in str(caught.value)
    assert KEY not in str(caught.value)


def test_connector_reports_no_internet_plainly():
    def opener(request, timeout):
        raise urllib.error.URLError("no route")

    with pytest.raises(MailError, match="internet"):
        ResendMailer(KEY, "minutes@example.com", opener=opener).send("a@b.co", "s", "h", "t")


def test_mailer_comes_only_from_the_environment():
    assert mailer_from_environment({}) is None
    assert mailer_from_environment({"RESEND_API_KEY": "  "}) is None
    with pytest.raises(MailConfigError):
        mailer_from_environment({"RESEND_API_KEY": KEY})
    with pytest.raises(MailConfigError):
        mailer_from_environment({"RESEND_API_KEY": KEY, "CONVENE_MAIL_FROM": "not an address"})
    for bad in ("<convene@example.com>", " <convene@example.com>", "Convene<convene@example.com>"):
        with pytest.raises(MailConfigError):  # Resend refuses these; startup must say so, not every send
            mailer_from_environment({"RESEND_API_KEY": KEY, "CONVENE_MAIL_FROM": bad})
    assert mailer_from_environment({"RESEND_API_KEY": KEY, "CONVENE_MAIL_FROM": "convene@example.com"}).sender == "convene@example.com"
    mailer = mailer_from_environment({"RESEND_API_KEY": KEY, "CONVENE_MAIL_FROM": "Convene <minutes@example.com>"})
    assert mailer.sender == "Convene <minutes@example.com>"


# --- the message --------------------------------------------------------------------------------


def test_message_escapes_meeting_values_and_names_the_attachment_safely():
    html, text = message.bodies("<b>Priya</b>", 'Q3 "plan" <script>', "30 September 2026")
    assert "<script>" not in html and "&lt;b&gt;Priya&lt;/b&gt;" in html
    assert "Hi <b>Priya</b>," in text and "30 September 2026" in text
    assert message.subject("Sprint planning") == "Your words, delivered: Sprint planning"
    assert message.attachment_name('Q3/Q4: "plan"') == "Convene minutes - Q3-Q4- -plan.docx"
    assert message.attachment_name("///") == "Convene minutes - meeting.docx"
    assert mask("priya@example.com") == "pr•••@example.com"


def headings(content: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(content)).paragraphs if p.style.name.startswith(("Heading", "Title"))]


def test_renderer_keeps_the_title_block_and_only_the_chosen_sections():
    from types import SimpleNamespace
    meeting = SimpleNamespace(title="Plan", meeting_id="m", started_at=None, created_at="2026-09-30T09:00:00.000Z")
    summary = SimpleNamespace(summary_text="We agreed.")
    args = (meeting, [], summary, [], [])
    names = lambda doc: [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert names(render(*args)) == ["Plan", "Summary", "Action items", "Transcript appendix"]
    assert names(render(*args, sections=("action_items",))) == ["Plan", "Action items"]
    assert names(render(*args, sections=("transcript", "summary"))) == ["Plan", "Summary", "Transcript appendix"]
    for bad in ((), ("decisions",)):
        with pytest.raises(ValueError):
            render(*args, sections=bad)


def test_message_lists_only_what_is_attached():
    assert message.contents(("action_items",)) == "the action items"
    assert message.contents(("summary", "action_items")) == "the summary and the action items"
    html, text = message.bodies("Maya", "Plan", "30 September 2026", ("summary", "action_items"))
    assert "attached: the summary and the action items." in text and "transcript" not in text
    assert "transcript" not in html


# --- the API --------------------------------------------------------------------------------------


@dataclass
class FakeMailer:
    sender: str = "Convene <minutes@example.com>"
    refuse: set = field(default_factory=set)
    sent: list = field(default_factory=list)

    def send(self, to, subject, html, text, attachment=None):
        if to in self.refuse:
            raise MailError("Resend HTTP 422: Invalid `to` field.")
        self.sent.append({"to": to, "subject": subject, "text": text, "attachment": attachment})
        return "msg"


@pytest.fixture
async def http():
    async with ClientSession() as session:
        yield session


async def serve(tmp_path, mailer):
    running = await start_server(settings_in(tmp_path), mailer=mailer)
    running.runtime.email.send_gap_s = 0
    return running


def with_emails(db, seeded, addresses):
    for name, email in addresses.items():
        participant_emails.set_email(db.conn, seeded.people[name], seeded.meeting_id, email, "2026-09-30T10:00:00.000Z")
    db.conn.commit()


@pytest.mark.asyncio
async def test_send_emails_each_recipient_separately_and_audits_ids_only(tmp_path, http):
    mailer = FakeMailer(refuse={"sam@example.org"})
    server = await serve(tmp_path, mailer)
    try:
        db = server.runtime.db
        seeded = meeting_at(db, "2026-09-30T09:00:00.000Z", title="Launch review", people=("Priya", "Sam", "Lee"))
        with_emails(db, seeded, {"Priya": "priya@example.com", "Sam": "sam@example.org"})
        url = f"{server.base_url}/api/meetings/{seeded.meeting_id}/email"
        async with http.get(url) as response:
            status = await response.json()
        assert status["configured"] and status["meeting_ended"]
        assert [r["email_masked"] for r in status["recipients"]] == ["pr•••@example.com", "sa•••@example.org"]
        assert "priya@example.com" not in json.dumps(status)

        async with http.post(url) as response:
            body = await response.json()
        assert response.status == 200
        assert body["sent_count"] == 1 and [f["display_name"] for f in body["failed"]] == ["Sam"]
        assert "Invalid" in body["failed"][0]["error"]
        sent = mailer.sent[0]
        assert sent["to"] == "priya@example.com" and sent["subject"] == "Your words, delivered: Launch review"
        assert sent["attachment"].filename == "Convene minutes - Launch review.docx"
        paragraphs = "\n".join(p.text for p in Document(io.BytesIO(sent["attachment"].content)).paragraphs)
        assert "Launch review" in paragraphs
        stamps = {r["display_name"]: r["last_sent_at"] for r in body["recipients"]}
        assert stamps["Priya"] is not None and stamps["Sam"] is None

        event = db.conn.execute("SELECT payload FROM AuditEvent WHERE event_type = 'minutes_emailed'").fetchone()[0]
        assert json.loads(event) == {"export_id": body["export_id"], "sent_participant_ids": [seeded.people["Priya"]],
                                     "failed_participant_ids": [seeded.people["Sam"]],
                                     "sections": {seeded.people[n]: ["summary", "action_items", "transcript"]
                                                  for n in ("Priya", "Sam")}}
        export_file = server.runtime.settings.exports_dir / f"{seeded.meeting_id}.docx"
        assert sent["attachment"].content == export_file.read_bytes()  # everything = the downloadable file itself
        everything = db.conn.execute("SELECT group_concat(payload) FROM AuditEvent").fetchone()[0]
        assert "@example" not in everything
    finally:
        await server.stop()


@pytest.mark.asyncio
async def test_send_refuses_until_it_can_send_and_sends_nothing(tmp_path, http):
    mailer = FakeMailer()
    server = await serve(tmp_path, mailer)
    try:
        db = server.runtime.db
        live = meeting_at(db, "2026-09-30T09:00:00.000Z", ended=False)
        with_emails(db, live, {"Priya": "priya@example.com"})
        nobody = meeting_at(db, "2026-09-30T09:00:00.000Z")
        unsummarized = meeting_at(db, "2026-09-30T09:00:00.000Z", summary="none")
        with_emails(db, unsummarized, {"Sam": "sam@example.com"})
        cases = [(live, "meeting_not_ended"), (nobody, "no_recipients"), (unsummarized, "summary_not_ready")]
        for seeded, code in cases:
            async with http.post(f"{server.base_url}/api/meetings/{seeded.meeting_id}/email") as response:
                assert (response.status, (await response.json())["error"]["code"]) == (409, code)
        async with http.post(f"{server.base_url}/api/meetings/{new_id()}/email") as response:
            assert response.status == 404
        assert mailer.sent == []
    finally:
        await server.stop()


@pytest.mark.asyncio
async def test_without_a_mailer_the_feature_says_it_is_not_configured(tmp_path, http):
    server = await serve(tmp_path, None)
    try:
        db = server.runtime.db
        seeded = meeting_at(db, "2026-09-30T09:00:00.000Z")
        with_emails(db, seeded, {"Priya": "priya@example.com"})
        url = f"{server.base_url}/api/meetings/{seeded.meeting_id}/email"
        async with http.get(url) as response:
            assert (await response.json())["configured"] is False
        async with http.post(url) as response:
            assert (response.status, (await response.json())["error"]["code"]) == (409, "mail_not_configured")
    finally:
        await server.stop()


@pytest.mark.asyncio
async def test_join_registration_accepts_email_and_never_returns_it(tmp_path, http):
    server = await serve(tmp_path, None)
    try:
        meeting = make_meeting(server.runtime.db)
        url = f"{server.base_url}/api/meetings/{meeting.meeting_id}/devices"
        payload = {"device_id": new_id(), "display_name": "Priya", "email": "priya@example.com"}
        async with http.post(url, json=payload) as response:
            text = await response.text()
        assert response.status == 201 and "priya@example.com" not in text
        async with http.post(url, json={**payload, "device_id": new_id(), "email": "nope"}) as response:
            assert (response.status, (await response.json())["error"]["code"]) == (400, "invalid_request")
        async with http.get(f"{server.base_url}/api/meetings/{meeting.meeting_id}") as response:
            assert "priya@example.com" not in await response.text()
    finally:
        await server.stop()


@pytest.mark.asyncio
async def test_each_person_gets_the_sections_chosen_for_them(tmp_path, http):
    mailer = FakeMailer()
    server = await serve(tmp_path, mailer)
    try:
        db = server.runtime.db
        seeded = meeting_at(db, "2026-09-30T09:00:00.000Z", title="Launch review", people=("Priya", "Sam", "Maya", "Lee"))
        with_emails(db, seeded, {"Priya": "priya@example.com", "Sam": "sam@example.com", "Maya": "maya@example.com",
                                 "Lee": "lee@example.com"})
        everything, short = ["transcript", "summary", "action_items"], ["action_items", "summary"]
        plan = [{"participant_id": seeded.people["Priya"], "sections": everything},
                {"participant_id": seeded.people["Sam"], "sections": everything},
                {"participant_id": seeded.people["Maya"], "sections": short}]
        async with http.post(f"{server.base_url}/api/meetings/{seeded.meeting_id}/email", json={"recipients": plan}) as response:
            body = await response.json()
        assert response.status == 200 and body["sent_count"] == 3
        got = {m["to"].split("@")[0]: m for m in mailer.sent}
        assert set(got) == {"priya", "sam", "maya"}  # Lee was not picked, so gets nothing
        assert headings(got["priya"]["attachment"].content) == ["Launch review", "Summary", "Action items", "Transcript appendix"]
        assert headings(got["maya"]["attachment"].content) == ["Launch review", "Summary", "Action items"]
        assert "the summary and the action items." in got["maya"]["text"]
        stamps = {r["display_name"]: r["last_sent_at"] for r in body["recipients"]}
        assert stamps["Lee"] is None and stamps["Maya"] is not None
        event = json.loads(db.conn.execute("SELECT payload FROM AuditEvent WHERE event_type = 'minutes_emailed'").fetchone()[0])
        assert event["sections"][seeded.people["Maya"]] == ["summary", "action_items"]  # stored in template order
        assert seeded.people["Lee"] not in event["sections"]
        assert db.conn.execute("SELECT COUNT(*) FROM Export WHERE status = 'ready'").fetchone()[0] == 1  # no extra exports
    finally:
        await server.stop()


@pytest.mark.asyncio
async def test_a_bad_selection_is_refused_before_anything_is_sent(tmp_path, http):
    mailer = FakeMailer()
    server = await serve(tmp_path, mailer)
    try:
        db = server.runtime.db
        seeded = meeting_at(db, "2026-09-30T09:00:00.000Z", people=("Priya", "Sam"))
        with_emails(db, seeded, {"Priya": "priya@example.com"})
        priya, sam = seeded.people["Priya"], seeded.people["Sam"]
        bad = [[], [{"participant_id": priya, "sections": []}], [{"participant_id": priya, "sections": ["decisions"]}],
               [{"participant_id": priya, "sections": ["summary"]}, {"participant_id": priya, "sections": ["summary"]}],
               [{"participant_id": sam, "sections": ["summary"]}],  # Sam gave no address
               [{"participant_id": new_id(), "sections": ["summary"]}], [{"sections": ["summary"]}]]
        for recipients in bad:
            async with http.post(f"{server.base_url}/api/meetings/{seeded.meeting_id}/email",
                                 json={"recipients": recipients}) as response:
                assert (response.status, (await response.json())["error"]["code"]) == (400, "invalid_request"), recipients
        assert mailer.sent == []
    finally:
        await server.stop()
