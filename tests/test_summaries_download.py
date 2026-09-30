"""GET /api/summaries/download: the history page's selected meetings' summaries as one DOCX."""
from dataclasses import replace
from io import BytesIO

import pytest
from aiohttp import ClientSession
from docx import Document

from server.config import ReportsConfig
from server.export.service import MIME
from server.ids import new_id
from server.rag.reasoning import FakeReasoningAdapter
from tests.support.reports import meeting_at
from tests.support.server import settings_in, start_server


@pytest.fixture
async def server(tmp_path):
    settings = replace(settings_in(tmp_path), reports=ReportsConfig(max_meetings=3))
    running = await start_server(settings, reasoning_adapter=FakeReasoningAdapter(
        fail=AssertionError("the summaries download must never call the reasoning model")))
    yield running
    await running.stop()


@pytest.fixture
async def http():
    async with ClientSession() as session:
        yield session


def counts(server):
    conn = server.runtime.db.conn
    return {t: conn.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0] for t in ("AuditEvent", "Export", "Summary")}


def url(server, *ids):
    return f"{server.base_url}/api/summaries/download?meeting_ids={','.join(ids)}"


@pytest.mark.asyncio
async def test_download_lists_each_meeting_oldest_first_with_its_summary(server, http):
    db = server.runtime.db
    late = meeting_at(db, "2026-09-25T10:00:00.000Z", title="Design review", items=(("Draw the icons", "Sam"),))
    early = meeting_at(db, "2026-09-22T10:00:00.000Z", title="Sprint planning")
    failed = meeting_at(db, "2026-09-23T10:00:00.000Z", title="Retro", summary="failed")
    none = meeting_at(db, "2026-09-24T10:00:00.000Z", title="Quick sync", summary="none")
    db.conn.execute("UPDATE Summary SET summary_text = ? WHERE meeting_id = ? AND status = 'ready'",
                    ("We ship Friday.\n\nPricing is still open.", early.meeting_id))
    db.conn.commit()
    before = counts(server)
    ids = [late.meeting_id, early.meeting_id, failed.meeting_id, late.meeting_id]  # a repeat is ignored
    async with http.get(url(server, *ids)) as response:
        content = await response.read()
        assert response.status == 200 and response.headers["Content-Type"] == MIME
        assert response.headers["Content-Disposition"].startswith('attachment; filename="convene-summaries-20')
    paragraphs = [p.text for p in Document(BytesIO(content)).paragraphs]
    headings = [p for p in paragraphs if " — 2026-" in p]
    assert headings == ["Sprint planning — 2026-09-22 UTC", "Retro — 2026-09-23 UTC", "Design review — 2026-09-25 UTC"]
    assert paragraphs[0] == "Convene meeting summaries"
    assert "3 meetings, 2 with a summary, oldest first (dates are UTC)" in paragraphs
    assert "We ship Friday." in paragraphs and "Pricing is still open." in paragraphs
    assert "No summary: the last attempt failed (summary_invalid_output)." in paragraphs
    assert "Quick sync" not in "\n".join(paragraphs)                 # not chosen
    assert "Draw the icons" not in "\n".join(paragraphs)             # summaries only, no action items
    assert counts(server) == before and server.runtime.reasoning_adapter.calls == []

    async with http.get(url(server, none.meeting_id)) as response:
        paragraphs = [p.text for p in Document(BytesIO(await response.read())).paragraphs]
    assert "No summary yet." in paragraphs and "1 meeting, 0 with a summary, oldest first (dates are UTC)" in paragraphs


@pytest.mark.asyncio
async def test_download_refuses_a_bad_selection(server, http):
    db = server.runtime.db
    ended = [meeting_at(db, "2026-09-22T10:00:00.000Z").meeting_id for _ in range(4)]
    live = meeting_at(db, "2026-09-22T10:00:00.000Z", ended=False).meeting_id
    cases = [(url(server), 400, "invalid_request"), (url(server, "not-a-uuid"), 400, "invalid_request"),
             (url(server, *ended), 400, "invalid_request"),        # over the cap of 3
             (url(server, new_id()), 404, "meeting_not_found"), (url(server, ended[0], live), 409, "meeting_not_ended")]
    for address, status, code in cases:
        async with http.get(address) as response:
            assert (response.status, (await response.json())["error"]["code"]) == (status, code), address
